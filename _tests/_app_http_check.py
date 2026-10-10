"""HTTP lifecycle check: REAL Flask app over a REAL socket, engine faked.

Starts app.py's Flask server on 127.0.0.1 (random free port) and drives the
actual HTTP surface with urllib:

  /            renders the UI
  /start       400 on empty target / garbage max; accepts a good job
  /status      running -> collecting -> done with full result payload
  /stop        sets the cooperative stop flag; job finishes "stopped"
  /download    real xlsx bytes (PK zip) that openpyxl re-opens with the rows
  409          second job while one is running (busy flag)
  404          unknown job / download before done

Run: .venv/Scripts/python.exe _tests/_app_http_check.py
"""
import faulthandler
import io
import json
import socket
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

faulthandler.dump_traceback_later(120, exit=True)

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import app as app_module  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

failures = []
def check(label, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + label + (f"  [{detail}]" if detail and not cond else ""),
          flush=True)
    if not cond:
        failures.append(label)


class FakeEngine:
    """Stand-in for scraper.scrape_followers honoring the real contract."""

    def __init__(self):
        self.calls = []
        self.hang = threading.Event()   # when set, the fake blocks until stop
        self.released = threading.Event()

    def __call__(self, target, headless=True, progress=None, max_followers=0,
                 should_stop=None, *args, **kwargs):
        self.calls.append({"target": target, "max": max_followers})
        names = [f"u{i}" for i in range(40)]
        for i in range(4):
            if should_stop and should_stop():
                break
            if self.hang.is_set():
                self.released.wait(timeout=20)
            time.sleep(0.15)
            if progress:
                progress((i + 1) * 10, expected=40, recent=names[: (i + 1) * 10])
        stopped = bool(should_stop and should_stop())
        return (names[: max_followers or None] if max_followers else names,
                False, stopped, 40, "stopped" if stopped else "exhausted", False)


fake = FakeEngine()
app_module.scrape_followers = fake  # patch the name _run_job actually calls
app_module._JOBS.clear()

with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
    s.bind(("127.0.0.1", 0))
    PORT = s.getsockname()[1]
BASE = f"http://127.0.0.1:{PORT}"

server = threading.Thread(
    target=lambda: app_module.app.run(host="127.0.0.1", port=PORT,
                                      debug=False, use_reloader=False,
                                      threaded=True),
    daemon=True)
server.start()

def http(path, data=None, want_status=None, method=None):
    if method is None:
        method = "POST" if data is not None else "GET"
    payload = urllib.parse.urlencode(data).encode() if data is not None else None
    req = urllib.request.Request(BASE + path, data=payload, method=method)
    try:
        with urllib.request.urlopen(req, timeout=10) as r:
            status, body, headers = r.status, r.read(), dict(r.headers)
    except urllib.error.HTTPError as e:
        status, body, headers = e.code, e.read(), dict(e.headers)
    if want_status is not None:
        check(f"{path} -> {want_status}", status == want_status, f"got {status}: {body[:120]}")
    return status, body, headers

for _ in range(50):  # wait for the server to answer
    try:
        http("/", want_status=200)
        break
    except Exception:
        time.sleep(0.2)

# --- UI renders -------------------------------------------------------------
_, body, _ = http("/", want_status=200)
check("index.html contains the Instagram UI with a /start fetch form",
      b"Instagram" in body and b"fetch(\"/start\"" in body and b"target" in body)

# --- app has no Instagram login endpoint -----------------------------------
http("/login-status", want_status=404)

# --- validation ---------------------------------------------------------------
http("/start", data={"target": ""}, want_status=400)
http("/start", data={"target": "nasa", "max": "5hundred"}, want_status=400)
http("/start", data={"target": "nasa", "max": "-5"}, want_status=400)

# --- happy path: start -> status -> done -> download -------------------------
status, body, _ = http("/start", data={"target": "nasa", "max": "40"}, want_status=200)
start_json = json.loads(body)
check("start returns a job id and mode", "job_id" in start_json and start_json["mode"] in {"anonymous", "logged_in"}, str(start_json))
job_id = start_json["job_id"]

seen_running = False
deadline = time.time() + 20
result = None
while time.time() < deadline:
    _, sbody, _ = http(f"/status/{job_id}", want_status=200)
    sj = json.loads(sbody)
    seen_running = seen_running or sj.get("status") == "running"
    if sj.get("status") in {"done", "error"}:
        result = sj
        break
    time.sleep(0.1)
check("job observed running before done", seen_running)
check("job reaches done", result is not None and result["status"] == "done", str(result))
res = (result or {}).get("result") or {}
check("result payload: scraped=40, expected=40, finish=exhausted",
      res.get("scraped") == 40 and res.get("expected") == 40
      and res.get("finish_reason") == "exhausted", str(res))
check("result payload: capped at the requested max (40)",
      fake.calls and fake.calls[-1]["max"] == 40, str(fake.calls))
check("result rows preview carries followers", len((res.get("rows") or {}).get("followers", [])) == 40)

status, xlsx_bytes, headers = http(f"/download/{job_id}", want_status=200)
check("download is an xlsx attachment",
      xlsx_bytes[:2] == b"PK" and "spreadsheet" in headers.get("Content-Type", ""),
      headers.get("Content-Type", ""))
check("download filename matches the report",
      (res.get("file") or "") in headers.get("Content-Disposition", ""),
      headers.get("Content-Disposition", ""))
wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
rows = list(wb["Followers"].iter_rows(values_only=True))
wb.close()
check("downloaded workbook has 40 follower rows",
      len(rows) == 41 and rows[1][1] == "u0" and rows[-1][1] == "u39",
      f"{len(rows)} rows")

# --- stop flow: a job interrupted by /stop finishes 'stopped' ----------------
fake.hang.set()
_, body, _ = http("/start", data={"target": "other"}, want_status=200)
stop_job = json.loads(body)["job_id"]
time.sleep(0.5)  # let the worker enter the hanging fake
http(f"/stop/{stop_job}", want_status=200, method="POST")
fake.released.set()
deadline = time.time() + 15
stopped_result = None
while time.time() < deadline:
    _, sbody, _ = http(f"/status/{stop_job}", want_status=200)
    sj = json.loads(sbody)
    if sj.get("status") in {"done", "error"}:
        stopped_result = sj
        break
    time.sleep(0.1)
check("stopped job ends done with stopped_early",
      stopped_result is not None and stopped_result["status"] == "done"
      and stopped_result["result"]["stopped_early"] is True
      and stopped_result["result"]["finish_reason"] == "stopped", str(stopped_result))
fake.hang.clear()

# --- 409: one scrape at a time -----------------------------------------------
fake.hang.set()
_, body, _ = http("/start", data={"target": "busytest"}, want_status=200)
busy_job = json.loads(body)["job_id"]
status, body, _ = http("/start", data={"target": "second"}, want_status=409)
check("409 payload flags busy", json.loads(body).get("busy") is True, body.decode())
http(f"/stop/{busy_job}", want_status=200, method="POST")
fake.released.set()
busy_result = None
deadline = time.time() + 15
while time.time() < deadline:
    _, sbody, _ = http(f"/status/{busy_job}", want_status=200)
    sj = json.loads(sbody)
    if sj.get("status") in {"done", "error"}:
        busy_result = sj
        break
    time.sleep(0.1)
check("stopped busy job also ends done", busy_result is not None and busy_result["status"] == "done")
fake.hang.clear()

# --- 404s ----------------------------------------------------------------------
http("/status/does-not-exist", want_status=404)
http("/download/does-not-exist", want_status=404)

# --- batch: several accounts -> one combined workbook ------------------------
http("/start", data={"target": "a, b", "names": "johndoe"}, want_status=400)
_, body, _ = http("/start", data={"target": "alpha, beta", "list": "following",
                                 "max": "10"}, want_status=200)
batch_start = json.loads(body)
check("start reports the account count and chosen list",
      batch_start.get("accounts") == 2 and batch_start.get("list") == "following",
      str(batch_start))
batch_job = batch_start["job_id"]
deadline = time.time() + 20
batch_result = None
while time.time() < deadline:
    _, sbody, _ = http(f"/status/{batch_job}", want_status=200)
    sj = json.loads(sbody)
    if sj.get("status") in {"done", "error"}:
        batch_result = sj
        break
    time.sleep(0.1)
check("batch job reaches done", batch_result is not None and batch_result["status"] == "done",
      str(batch_result))
bres = (batch_result or {}).get("result") or {}
accounts = bres.get("accounts") or []
check("batch result lists every account with counts",
      bres.get("mode") == "batch_scrape" and bres.get("list_type") == "following"
      and [a["account"] for a in accounts] == ["alpha", "beta"]
      and all(a["collected"] == 10 for a in accounts), str(bres)[:300])
status, xlsx_bytes, _ = http(f"/download/{batch_job}", want_status=200)
wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
sheets = wb.sheetnames
summary = list(wb["Summary"].iter_rows(values_only=True))
wb.close()
check("combined workbook: Summary + one sheet per account",
      sheets == ["Summary", "alpha", "beta"], str(sheets))
check("summary sheet has one row per account",
      len(summary) == 3 and summary[1][1] == "@alpha" and summary[2][1] == "@beta"
      and summary[1][2] == "following" and summary[1][3] == 10, str(summary))

# --- retry: a transient failure gets exactly one end-of-batch retry ----------
from unittest import mock  # noqa: E402

retry_calls = []

def flaky(target, headless=True, progress=None, max_followers=0,
          should_stop=None, list_type="followers"):
    retry_calls.append(target)
    if target == "beta" and retry_calls.count("beta") == 1:
        raise app_module.ScrapeError("transient hiccup")
    return ([f"{target}1", f"{target}2"], False, False, 2, "exhausted", False)

saved_engine = app_module.scrape_followers
app_module.scrape_followers = flaky
retry_result = None
try:
    with mock.patch.object(app_module, "RETRY_DELAY_SECONDS", 0):
        _, body, _ = http("/start", data={"target": "alpha, beta",
                                          "list": "following"}, want_status=200)
        retry_job = json.loads(body)["job_id"]
        deadline = time.time() + 20
        while time.time() < deadline:
            _, sbody, _ = http(f"/status/{retry_job}", want_status=200)
            sj = json.loads(sbody)
            if sj.get("status") in {"done", "error"}:
                retry_result = sj
                break
            time.sleep(0.1)
finally:
    app_module.scrape_followers = saved_engine

check("retry: job completes despite the transient failure",
      retry_result is not None and retry_result["status"] == "done", str(retry_result))
rres = (retry_result or {}).get("result") or {}
check("retry: the failed account was attempted exactly twice",
      retry_calls == ["alpha", "beta", "beta"], str(retry_calls))
check("retry: both accounts end ok",
      {a["account"]: a["status"] for a in rres.get("accounts", [])}
      == {"alpha": "ok", "beta": "ok"}, str(rres)[:200])
status, xlsx_bytes, _ = http(f"/download/{retry_job}", want_status=200)
wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
retry_sheets = wb.sheetnames
wb.close()
check("retry: workbook carries both recovered accounts",
      retry_sheets == ["Summary", "alpha", "beta"], str(retry_sheets))

# cleanup: remove the xlsx files this check created in output/
for job in (result, stopped_result, busy_result, batch_result, retry_result):
    name = ((job or {}).get("result") or {}).get("file")
    if name:
        (ROOT / "output" / name).unlink(missing_ok=True)

print("ALL PASS" if not failures else f"{len(failures)} CHECK(S) FAILED", flush=True)
sys.exit(1 if failures else 0)
