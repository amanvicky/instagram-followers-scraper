"""Scraper engine tests (pure logic + Playwright-page fakes — no browser).

Small tests: response-shape parsing, username filtering, dedup math, Excel
metadata. One RED pair pins the known graphql-extractor bug (nested
node.user.username payloads return no names). Run from the repo root:
    python -m unittest discover -s tests -t . -v
"""
import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import scraper  # noqa: E402


# --- pure logic: reference-list parsing (user-facing compare feature) --------

class ReferenceParsingTests(unittest.TestCase):
    def test_extracts_at_mentions_urls_and_plain_names(self):
        raw = ["@jane_smith", "https://www.instagram.com/someone/",
               "johndoe", "John Doe (@mixed.case)"]
        names = scraper.normalize_usernames(raw)
        self.assertEqual(names, ["jane_smith", "someone", "johndoe",
                                 "mixed.case"])

    def test_deduplicates_case_insensitively_and_keeps_first_spelling(self):
        names = scraper.normalize_usernames(["JohnDoe", "johndoe", "JOHNDOE"])
        self.assertEqual(names, ["JohnDoe"])

    def test_rejects_reserved_segments_and_numbers(self):
        names = scraper.normalize_usernames(["followers", "12345", "p",
                                             "real_user"])
        self.assertEqual(names, ["real_user"])


# --- pure logic: pagination response shapes ----------------------------------

class ResponseShapeTests(unittest.TestCase):
    def test_graphql_edges_shape_extracts_names_and_cursor(self):
        payload = {"data": {"user": {"edges": [
            {"node": {"username": "a"}}, {"node": {"username": "b"}}],
            "page_info": {"end_cursor": "CUR", "has_next_page": True}}}}
        names, cursor, has_more = scraper._extract_names_and_cursor(payload)
        self.assertEqual(names, ["a", "b"])
        self.assertEqual(cursor, "CUR")
        self.assertTrue(has_more)

    def test_friendships_shape_extracts_names_and_max_id(self):
        payload = {"users": [{"username": "x"}, {"username": "y"}],
                   "next_max_id": "NEXT", "has_more": True}
        names, cursor, has_more = scraper._extract_names_and_cursor(payload)
        self.assertEqual(names, ["x", "y"])
        self.assertEqual(cursor, "NEXT")
        self.assertTrue(has_more)

    # RED pair: Instagram ships a shape where the user object is nested one
    # level deeper (node.user.username). The current extractor returns NO
    # names for it, so the API fallback would stall at zero. These tests fail
    # until _extract_names_and_cursor handles the nested shape.

    def test_nested_graphql_user_shape_extracts_names(self):
        payload = {"data": {"user": {"edge_followed_by": {"edges": [
            {"node": {"user": {"username": "deep_one"}}},
            {"node": {"user": {"username": "deep_two"}}}],
            "page_info": {"end_cursor": "D", "has_next_page": False}}}}}
        names, cursor, has_more = scraper._extract_names_and_cursor(payload)
        self.assertEqual(names, ["deep_one", "deep_two"])
        self.assertEqual(cursor, "D")
        self.assertFalse(has_more)

    def test_nested_shape_with_junk_rows_skips_nulls(self):
        payload = {"edges": [{"node": None},
                             {"node": {"user": None}},
                             {"node": {"user": {"username": "ok"}}}],
                   "page_info": {"end_cursor": None, "has_next_page": False}}
        names, _, _ = scraper._extract_names_and_cursor(payload)
        self.assertEqual(names, ["ok"])


# --- collection loop: dedup + counters, browser fully faked -------------------

def _dialog_page(rows_per_round, rounds_before_stall=2):
    """Fake page: returns the same visible rows every round (a stalled
    dialog), never grows, never opens a login wall."""
    page = mock.Mock()
    page.url = "https://www.instagram.com/nasa/followers/"
    state = {"rounds": 0}
    page.evaluate.side_effect = _evaluate_fake(rows_per_round, state)
    return page


def _evaluate_fake(rows_per_round, state):
    def fake_evaluate(script, *args):
        if "_COLLECT_JS" in script or "querySelectorAll('a[href]')" in script:
            state["rounds"] += 1
            return [f"row{i}" for i in range(rows_per_round)]
        if "_LOGIN_WALL_JS" in script:
            return False
        if "_SCROLL_JS" in script:
            return True
        if "csrftoken" in script:
            return "TOKEN"
        return None
    return fake_evaluate


class CollectLoopDedupTests(unittest.TestCase):
    def test_repeated_visible_rows_are_counted_once(self):
        page = _dialog_page(rows_per_round=5)
        with mock.patch.object(scraper, "MAX_API_ATTEMPTS", 0), \
             mock.patch.object(scraper, "SCROLL_RETRIES", 3):
            followers, wall, early, expected, reason, suggested = \
                scraper.collect_followers(page, "nasa", timeout=5)
        self.assertEqual(len(followers), 5)   # 5 rows x 1, not 5 x rounds
        self.assertEqual(len(set(followers)), len(followers))

    def test_finish_reason_is_exhausted_when_list_settles(self):
        page = _dialog_page(rows_per_round=5)
        with mock.patch.object(scraper, "MAX_API_ATTEMPTS", 0), \
             mock.patch.object(scraper, "SCROLL_RETRIES", 3):
            *_, reason, _suggested = scraper.collect_followers(page, "nasa", timeout=5)
        self.assertEqual(reason, "exhausted")


class LoginWallHandlingTests(unittest.TestCase):
    def test_login_wall_returns_empty_public_result_without_exception(self):
        page = mock.Mock()
        page.url = "https://www.instagram.com/nasa/followers/"

        def evaluate(script, *args):
            if "document.body ? document.body.innerText" in script:
                return ""
            if "log in to see|log in to view|sign up to see" in script:
                return True
            if "querySelectorAll('a[href]')" in script:
                return []
            if "split('\\n')" in script:
                return []
            return None

        page.evaluate.side_effect = evaluate
        with mock.patch.object(scraper, "MAX_API_ATTEMPTS", 0):
            followers, wall, stopped, expected, reason, suggestions = \
                scraper.collect_followers(page, "nasa", timeout=1)

        self.assertEqual(followers, [])
        self.assertTrue(wall)
        self.assertFalse(stopped)
        self.assertIsNone(expected)
        self.assertEqual(reason, "login_wall")
        self.assertFalse(suggestions)


class BrowserSessionTests(unittest.TestCase):
    def test_scrape_uses_ephemeral_context_without_saved_profile(self):
        playwright = mock.MagicMock()
        browser = playwright.chromium.launch.return_value
        context = browser.new_context.return_value
        context.pages = [mock.Mock()]
        runner = mock.MagicMock()
        runner.__enter__.return_value = playwright
        with mock.patch.object(scraper, "sync_playwright", return_value=runner), \
             mock.patch.object(scraper, "collect_followers",
                               return_value=([], False, False, None,
                                             "exhausted", False)):
            scraper.scrape_followers("nasa", headless=True)

        playwright.chromium.launch.assert_called_once_with(headless=True)
        browser.new_context.assert_called_once()
        playwright.chromium.launch_persistent_context.assert_not_called()
        context.close.assert_called_once()
        browser.close.assert_called_once()


# --- following-list suggested-accounts detection ------------------------------

class FollowingSuggestionTests(unittest.TestCase):
    def test_follower_lists_are_never_flagged(self):
        cleaned, flagged = scraper._flag_following_suggestions(
            ["a", "b"], expected=1, list_type="followers")
        self.assertFalse(flagged)
        self.assertEqual(cleaned, ["a", "b"])

    def test_padded_following_list_is_flagged(self):
        # Verified live: counter says 1, dialog/API served 33 names.
        cleaned, flagged = scraper._flag_following_suggestions(
            [f"u{i}" for i in range(33)], expected=1, list_type="following")
        self.assertTrue(flagged)

    def test_normal_following_list_is_not_flagged(self):
        cleaned, flagged = scraper._flag_following_suggestions(
            [f"u{i}" for i in range(30)], expected=89, list_type="following")
        self.assertFalse(flagged)

    def test_tiny_list_over_a_tiny_counter_is_not_flagged(self):
        # A 5-name collection over a counter of 1 is plausible reading noise,
        # not proof of padding - the threshold is max(10, counter).
        _, flagged = scraper._flag_following_suggestions(
            ["u1", "u2", "u3", "u4", "u5"], expected=1, list_type="following")
        self.assertFalse(flagged)


# --- Excel metadata: the honest-completeness feature --------------------------

class ExcelMetadataTests(unittest.TestCase):
    def _save(self, tmp, scraped, expected, finish_reason, partial=False):
        return scraper.save_to_excel(
            "target", [f"u{i}" for i in range(scraped)],
            dt.datetime(2026, 9, 28, 12, 0, 0), Path(tmp),
            partial, expected=expected, finish_reason=finish_reason)

    def _info_map(self, path):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        ws = wb["Info"]
        rows = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r[0]}
        wb.close()
        return rows

    def test_complete_run_says_appears_complete(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._save(tmp, scraped=25, expected=25,
                              finish_reason="exhausted")
            info = self._info_map(path)
        self.assertEqual(info["List complete?"], "Appears complete")
        self.assertEqual(info["Instagram's follower counter"], 25)
        self.assertEqual(info["Gap (counter minus collected)"], 0)

    def test_throttled_run_with_gap_flags_possibly_incomplete(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._save(tmp, scraped=20, expected=25,
                              finish_reason="throttled")
            info = self._info_map(path)
        self.assertIn("Possibly incomplete", info["List complete?"])
        self.assertEqual(info["Gap (counter minus collected)"], 5)


# --- pure logic: run-to-run diff (convergence measurement) -------------------

class DiffFollowersTests(unittest.TestCase):
    # Contract: diff_followers(NEW, OLD) -> (names new since OLD,
    # names OLD had that are now gone) — matches the CLI's --diff NEW OLD.
    def test_new_names_are_those_not_in_the_old_run(self):
        new_names, lost = scraper.diff_followers(
            ["b", "c", "d", "e"], ["a", "b", "c"])
        self.assertEqual(new_names, ["d", "e"])

    def test_lost_names_are_those_missing_from_the_new_run(self):
        new_names, lost = scraper.diff_followers(
            ["b", "c", "d"], ["a", "b", "c"])
        self.assertEqual(lost, ["a"])

    def test_diff_is_case_insensitive_and_keeps_first_spelling(self):
        new_names, lost = scraper.diff_followers(
            ["johndoe", "new_one"], ["JohnDoe"])
        self.assertEqual(new_names, ["new_one"])
        self.assertEqual(lost, [])

    def test_identical_runs_yield_no_new_and_no_lost(self):
        names = ["u1", "u2", "u3"]
        new_names, lost = scraper.diff_followers(names, list(names))
        self.assertEqual((new_names, lost), ([], []))


if __name__ == "__main__":
    unittest.main()
