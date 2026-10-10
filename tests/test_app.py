"""Flask app tests (stdlib unittest + Flask test client).

Medium-size tests: real app routes, faked engine — the browser boundary
(``app.scrape_followers``) is patched, so nothing touches Instagram and no
Chromium is launched. Run from the repo root:
    python -m unittest discover -s tests -t . -v
"""
import sys
import threading
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import app as app_module  # noqa: E402
import scraper as scraper_module  # noqa: E402


def _fake_scrape(scraped=20, expected=25, finish_reason="throttled"):
    """A controlled engine result: the six-tuple _run_job unpacks."""
    followers = [f"user{i}" for i in range(scraped)]
    return (followers, False, False, expected, finish_reason, False)


class StartRouteTests(unittest.TestCase):
    def setUp(self):
        app_module.app.config["TESTING"] = True
        # Tests that stub the worker thread leave their job record behind;
        # clear it so no leftover "running" job trips the 409 guard below.
        app_module._JOBS.clear()
        self.client = app_module.app.test_client()

    def tearDown(self):
        app_module._JOBS.clear()

    # --- existing behavior (pin) -------------------------------------------

    def test_start_rejects_empty_target_with_400(self):
        res = self.client.post("/start", data={"target": ""})
        self.assertEqual(res.status_code, 400)
        self.assertIn("target", res.get_json()["error"].lower())

    def test_start_counts_pasted_reference_names(self):
        with mock.patch.object(app_module.threading, "Thread") as thread:
            thread.return_value.start = lambda: None  # never run the worker
            res = self.client.post("/start", data={
                "target": "nasa",
                "names": "johndoe\n@jane_x\nhttps://www.instagram.com/someone/",
            })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.get_json()["reference_size"], 3)

    # --- RED: server-side cap validation (bug: garbage max runs uncapped) ---
    # The UI rejects these; the server silently coerces them to 0 (= uncapped).
    # These two tests fail until /start enforces the same digits-only rule.

    def test_start_rejects_non_numeric_max_with_400(self):
        res = self.client.post("/start", data={"target": "nasa", "max": "5hundred"})
        self.assertEqual(res.status_code, 400,
                         "garbage max must be rejected, not run uncapped")
        self.assertIn("number", res.get_json()["error"].lower())

    def test_start_rejects_negative_max_with_400(self):
        res = self.client.post("/start", data={"target": "nasa", "max": "-5"})
        self.assertEqual(res.status_code, 400,
                         "negative max must be rejected, not run uncapped")

    # --- multi-account parsing -------------------------------------------------

    def test_parse_targets_splits_commas_whitespace_and_dedups(self):
        self.assertEqual(
            app_module.parse_targets("nasa, nintendo  adi,@space_odds"),
            ["nasa", "nintendo", "adi", "space_odds"])

    def test_parse_targets_drops_junk_and_keeps_first_spelling(self):
        self.assertEqual(app_module.parse_targets("a, @A, , ,,b"), ["a", "b"])

    def test_parse_targets_does_not_silently_discard_entries_after_limit(self):
        raw = ", ".join(f"user{i}" for i in range(app_module.MAX_TARGETS + 1))
        self.assertEqual(len(app_module.parse_targets(raw)), app_module.MAX_TARGETS + 1)

    def test_start_rejects_target_count_over_limit(self):
        raw = ", ".join(f"user{i}" for i in range(app_module.MAX_TARGETS + 1))
        res = self.client.post("/start", data={"target": raw})
        self.assertEqual(res.status_code, 400)
        self.assertIn(str(app_module.MAX_TARGETS), res.get_json()["error"])

    def test_compare_mode_always_scrapes_followers(self):
        with mock.patch.object(app_module, "scrape_followers",
                               return_value=_fake_scrape()) as fake, \
             mock.patch.object(app_module.threading, "Thread") as thread:
            thread.side_effect = lambda target, args, daemon: (
                mock.Mock(start=lambda: target(*args)))
            res = self.client.post("/start", data={
                "target": "nasa", "names": "user0", "list": "following",
            })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(fake.call_args.kwargs["list_type"], "followers")

    # --- engine handoff contract ---------------------------------------------

    def test_start_passes_numeric_max_to_the_engine(self):
        with mock.patch.object(app_module, "scrape_followers",
                               return_value=_fake_scrape()) as fake, \
             mock.patch.object(app_module.threading, "Thread") as thread:
            thread.side_effect = lambda target, args, daemon: (
                mock.Mock(start=lambda: target(*args)))
            res = self.client.post("/start", data={"target": "nasa", "max": "500"})
            self.assertEqual(res.status_code, 200)
            thread.return_value.start()
        kwargs = fake.call_args.kwargs
        self.assertEqual(kwargs.get("max_followers"), 500)
        self.assertTrue(callable(kwargs.get("should_stop")))


class RunJobTests(unittest.TestCase):
    """_run_job called directly (synchronous) with the engine faked.

    In production /start registers the job in _JOBS before threading the
    worker, so every test here registers its job the same way first.
    """

    def setUp(self):
        app_module.app.config["TESTING"] = True
        self.job_id = "testjob"
        self._register_job(self.job_id)
        self.addCleanup(app_module._JOBS.pop, self.job_id, None)

    @staticmethod
    def _register_job(job_id):
        app_module._JOBS[job_id] = {
            "status": "running", "phase": "Starting browser", "count": 0,
            "expected": None, "recent": [], "result": None, "message": None,
        }

    def test_done_job_carries_scraped_expected_and_finish_reason(self):
        with mock.patch.object(app_module, "scrape_followers",
                               return_value=_fake_scrape(20, 25, "throttled")):
            app_module._run_job(self.job_id, ["nasa"], [], True, 0,
                                "followers", None)
        job = app_module._JOBS[self.job_id]
        self.assertEqual(job["status"], "done")
        self.assertEqual(job["result"]["scraped"], 20)
        self.assertEqual(job["result"]["expected"], 25)
        self.assertEqual(job["result"]["finish_reason"], "throttled")
        self.assertEqual(job["result"]["file"].endswith(".xlsx"), True)

    def test_public_login_wall_with_zero_names_saves_report_without_login(self):
        import tempfile
        from pathlib import Path as P
        from openpyxl import load_workbook

        empty = ([], True, False, 25, "login_wall", False)
        output_dir = P(tempfile.mkdtemp())
        with mock.patch.object(app_module, "scrape_followers", return_value=empty), \
             mock.patch.object(app_module, "OUTPUT_DIR", output_dir):
            app_module._run_job(self.job_id, ["nasa"], [], True, 0,
                                "followers", None)

        job = app_module._JOBS[self.job_id]
        self.assertEqual(job["status"], "done")
        self.assertEqual(job["result"]["scraped"], 0)
        self.assertEqual(job["result"]["finish_reason"], "login_wall")
        workbook = load_workbook(output_dir / job["result"]["file"], read_only=True)
        info = {row[0]: row[1] for row in workbook["Info"].iter_rows(values_only=True)
                if row[0]}
        workbook.close()
        self.assertEqual(info["Mode"], "Anonymous (no login)")
        self.assertIn("login wall", info["List complete?"].lower())

    def test_engine_receives_max_followers_and_stop_flag(self):
        with mock.patch.object(app_module, "scrape_followers",
                               return_value=_fake_scrape()) as fake:
            stop = threading.Event()
            app_module._run_job(self.job_id, ["nasa"], [], True, 750,
                                "followers", stop)
        kwargs = fake.call_args.kwargs
        self.assertEqual(kwargs.get("max_followers"), 750)
        self.assertEqual(kwargs.get("list_type"), "followers")
        self.assertTrue(callable(kwargs.get("should_stop")))


def _batch_scrape(one=3, two=2, one_expected=3, two_expected=None,
                  fail_second=False):
    """Fake engine returning different sizes per account, optionally failing."""
    def fake(target, headless=True, progress=None, max_followers=0,
             should_stop=None, list_type="followers"):
        if fail_second and target == "b":
            raise scraper_module.ScrapeError("private account")
        names = [f"{target}{i}" for i in range(one if target == "a" else two)]
        expected = one_expected if target == "a" else two_expected
        return (names, False, False, expected, "exhausted", False)
    return fake


class BatchRunJobTests(unittest.TestCase):
    """Multi-account batch runs: sequential per-account collection with
    per-account error isolation and one combined workbook."""

    def setUp(self):
        app_module.app.config["TESTING"] = True
        self.job_id = "batchjob"
        app_module._JOBS[self.job_id] = {
            "status": "running", "phase": "Starting browser", "count": 0,
            "expected": None, "recent": [], "result": None, "message": None,
        }
        self.addCleanup(app_module._JOBS.pop, self.job_id, None)

    def test_batch_writes_one_workbook_with_per_account_sheets(self):
        import tempfile
        from pathlib import Path as P
        from openpyxl import load_workbook
        with mock.patch.object(app_module, "scrape_followers",
                               side_effect=_batch_scrape()), \
             mock.patch.object(app_module, "OUTPUT_DIR", P(tempfile.mkdtemp())):
            app_module._run_job(self.job_id, ["a", "b"], [], True, 0,
                                "following", None)
            wb_path = app_module.OUTPUT_DIR / app_module._JOBS[self.job_id]["result"]["file"]
        job = app_module._JOBS[self.job_id]
        self.assertEqual(job["status"], "done")
        result = job["result"]
        self.assertEqual(result["mode"], "batch_scrape")
        self.assertEqual(result["list_type"], "following")
        self.assertEqual([a["account"] for a in result["accounts"]], ["a", "b"])
        self.assertEqual(result["accounts"][0]["collected"], 3)
        self.assertEqual(result["accounts"][1]["collected"], 2)
        wb = load_workbook(wb_path, read_only=True)
        self.assertIn("Summary", wb.sheetnames)
        self.assertIn("a", wb.sheetnames)
        self.assertIn("b", wb.sheetnames)
        wb.close()

    def test_batch_keeps_going_when_one_account_fails(self):
        import tempfile
        from pathlib import Path as P
        with mock.patch.object(app_module, "scrape_followers",
                               side_effect=_batch_scrape(fail_second=True)), \
             mock.patch.object(app_module, "RETRY_DELAY_SECONDS", 0), \
             mock.patch.object(app_module, "OUTPUT_DIR", P(tempfile.mkdtemp())):
            app_module._run_job(self.job_id, ["a", "b"], [], True, 0,
                                "following", None)
        job = app_module._JOBS[self.job_id]
        self.assertEqual(job["status"], "done")   # not an overall error
        result = job["result"]
        statuses = {a["account"]: a["status"] for a in result["accounts"]}
        self.assertEqual(statuses, {"a": "ok", "b": "failed"})
        self.assertIn("private account", result["accounts"][1]["error"])
        self.assertEqual(result["done"], 1)
        self.assertEqual(result["failed"], 1)

    def test_stop_flag_prevents_later_accounts_from_starting(self):
        seen = []
        def fake(target, **kwargs):
            seen.append(target)
            return ([f"{target}1"], False, False, 1, "exhausted", False)
        with mock.patch.object(app_module, "scrape_followers",
                               side_effect=fake):
            stop = threading.Event()
            stop.set()   # already stopped before the first account
            app_module._run_job(self.job_id, ["a", "b"], [], True, 0,
                                "following", stop)
        self.assertEqual(seen, [])   # no account started
        job = app_module._JOBS[self.job_id]
        self.assertEqual(job["status"], "error")   # nothing collected at all

    def test_failed_account_is_retried_once_and_recovered(self):
        import tempfile
        from pathlib import Path as P
        calls = []

        def flaky(target, **kwargs):
            calls.append(target)
            if len(calls) == 1:   # only the very first attempt fails
                raise scraper_module.ScrapeError("transient hiccup")
            return ([f"{target}1", f"{target}2"], False, False, 2, "exhausted", False)

        with mock.patch.object(app_module, "scrape_followers",
                               side_effect=flaky), \
             mock.patch.object(app_module, "RETRY_DELAY_SECONDS", 0), \
             mock.patch.object(app_module, "OUTPUT_DIR", P(tempfile.mkdtemp())):
            app_module._run_job(self.job_id, ["a", "b"], [], True, 0,
                                "following", None)
        job = app_module._JOBS[self.job_id]
        self.assertEqual(job["status"], "done")
        statuses = {a["account"]: a["status"] for a in job["result"]["accounts"]}
        self.assertEqual(statuses, {"a": "ok", "b": "ok"})
        self.assertEqual(calls, ["a", "b", "a"])   # a retried after the main pass

    def test_single_account_failure_is_not_retried(self):
        with mock.patch.object(app_module, "scrape_followers",
                               side_effect=_batch_scrape(fail_second=True)) as fake:
            app_module._run_job(self.job_id, ["b"], [], True, 0,
                                "following", None)
        self.assertEqual(fake.call_count, 1)   # one attempt, no retry pass
        job = app_module._JOBS[self.job_id]
        self.assertEqual(job["status"], "error")

    def test_multi_target_with_reference_names_is_rejected_by_start(self):
        app_module.app.config["TESTING"] = True
        client = app_module.app.test_client()
        app_module._JOBS.clear()
        res = client.post("/start", data={
            "target": "a, b",
            "names": "johndoe\njane",
        })
        self.assertEqual(res.status_code, 400)
        self.assertIn("single", res.get_json()["error"].lower())


class StatusAndDownloadTests(unittest.TestCase):
    def setUp(self):
        app_module.app.config["TESTING"] = True
        self.client = app_module.app.test_client()

    def test_status_of_unknown_job_is_404(self):
        self.assertEqual(self.client.get("/status/does-not-exist").status_code, 404)

    def test_public_app_does_not_expose_instagram_login_routes(self):
        self.assertEqual(self.client.post("/login").status_code, 404)
        self.assertEqual(self.client.get("/login-status").status_code, 404)

    def test_download_of_unknown_job_is_404(self):
        self.assertEqual(
            self.client.get("/download/does-not-exist").status_code, 404)

    def test_download_before_done_is_404(self):
        app_module._JOBS["partial-job"] = {"status": "running", "result": None}
        try:
            self.assertEqual(
                self.client.get("/download/partial-job").status_code, 404)
        finally:
            app_module._JOBS.pop("partial-job", None)


if __name__ == "__main__":
    unittest.main()
